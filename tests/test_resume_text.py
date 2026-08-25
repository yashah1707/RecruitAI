"""Tests for PDF/DOCX text extraction, including the extractor fallback chain."""

from __future__ import annotations

import io

import pymupdf
import pytest
from docx import Document

from app.resume_text import MIN_TEXT_CHARS, UnreadableResumeError, extract_text


def _pdf_bytes(lines: list[str]) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "\n".join(lines), fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def _docx_bytes(lines: list[str]) -> bytes:
    d = Document()
    for line in lines:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


SYNTHETIC = [
    "A. Synthetic Candidate",
    "PhD Scholar in Computer Science and Engineering",
    "M.Sc. Physics, Synthetic University, 2015",
]


def test_pdf_text_is_extracted():
    r = extract_text(io.BytesIO(_pdf_bytes(SYNTHETIC)), "cv.pdf")
    assert "Synthetic Candidate" in r.text
    assert r.source_filename == "cv.pdf"


def test_pdf_words_are_not_glued_together():
    """The bug this switch fixed: some layouts lost inter-word spaces, giving
    'PhDScholarinComputerScience...' which is unreadable in the evidence
    column a human reviewer relies on."""
    r = extract_text(io.BytesIO(_pdf_bytes(SYNTHETIC)), "cv.pdf")
    assert "PhD Scholar in Computer Science" in r.text


def test_docx_text_is_extracted():
    r = extract_text(io.BytesIO(_docx_bytes(SYNTHETIC)), "cv.docx")
    assert "Synthetic Candidate" in r.text


def test_docx_table_cells_are_included():
    d = Document()
    d.add_paragraph("A. Synthetic Candidate")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "UGC-NET"
    t.rows[0].cells[1].text = "December 2016"
    buf = io.BytesIO()
    d.save(buf)
    r = extract_text(io.BytesIO(buf.getvalue()), "cv.docx")
    assert "UGC-NET" in r.text and "December 2016" in r.text


def test_corrupted_pdf_fails_cleanly_without_crashing():
    with pytest.raises(UnreadableResumeError):
        extract_text(io.BytesIO(b"%PDF-1.4 truncated garbage"), "broken.pdf")


def test_pdf_with_no_text_layer_is_reported_as_unreadable():
    """A scanned image PDF must be a clean failure, never a silent empty
    extraction that then gets sent to the LLM."""
    doc = pymupdf.open()
    doc.new_page()  # blank page, no text
    data = doc.tobytes()
    doc.close()
    with pytest.raises(UnreadableResumeError, match="no text layer|could not read"):
        extract_text(io.BytesIO(data), "scanned.pdf")


def test_near_empty_text_is_rejected():
    r = _pdf_bytes(["hi"])
    with pytest.raises(UnreadableResumeError):
        extract_text(io.BytesIO(r), "tiny.pdf")


def test_unsupported_extension_is_rejected():
    with pytest.raises(UnreadableResumeError, match="unsupported file type"):
        extract_text(io.BytesIO(b"whatever"), "notes.txt")


def test_falls_back_to_pdfplumber_when_pymupdf_fails(monkeypatch):
    """The two libraries fail on different malformed files, so a failure in
    the primary must not lose a file the fallback could have read."""
    import app.resume_text as rt

    def boom(fileobj):
        raise RuntimeError("simulated pymupdf failure")

    monkeypatch.setattr(rt, "_extract_pdf_pymupdf", boom)
    r = extract_text(io.BytesIO(_pdf_bytes(SYNTHETIC)), "cv.pdf")
    assert "Synthetic Candidate" in r.text


def test_error_names_both_extractors_when_both_fail(monkeypatch):
    import app.resume_text as rt

    monkeypatch.setattr(rt, "_extract_pdf_pymupdf", lambda f: (_ for _ in ()).throw(RuntimeError("boom-a")))
    monkeypatch.setattr(rt, "_extract_pdf_pdfplumber", lambda f: (_ for _ in ()).throw(RuntimeError("boom-b")))
    with pytest.raises(UnreadableResumeError) as exc:
        extract_text(io.BytesIO(_pdf_bytes(SYNTHETIC)), "cv.pdf")
    assert "pymupdf" in str(exc.value) and "pdfplumber" in str(exc.value)
