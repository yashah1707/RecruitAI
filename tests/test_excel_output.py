"""FakeProvider -> excel_writer, verified by reopening the written workbook.

No model runs in this file.
"""

from __future__ import annotations

from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from app.excel_writer import (
    COLUMNS,
    EVIDENCE_COLUMNS,
    SHEET_NAME,
    build_row,
    evidence_column,
    export_name,
    write_workbook,
)
from llm.interface import FIELD_NAMES, ResumeRecord
from llm.providers.fake_provider import FakeProvider

THRESHOLD = 0.7


@pytest.fixture
def records() -> list[ResumeRecord]:
    """Three fake extractions plus one file that never reached the model."""
    provider = FakeProvider()
    recs = [
        ResumeRecord(
            source_filename=f"synthetic_{i}.pdf",
            processed_at=datetime(2026, 8, 22, 10, 30, i),
            result=provider.extract_fields("resume text"),
        )
        for i in range(3)
    ]
    recs.append(ResumeRecord.failed("scanned_image.pdf", "no extractable text in file"))
    return recs


@pytest.fixture
def sheet(records, tmp_path):
    path = write_workbook(records, tmp_path / "out.xlsx", threshold=THRESHOLD)
    return load_workbook(path)[SHEET_NAME]


def header(sheet) -> list[str]:
    return [c.value for c in sheet[1]]


def column_index(sheet, name: str) -> int:
    return header(sheet).index(name) + 1


def test_header_matches_the_declared_column_order(sheet):
    assert header(sheet) == list(COLUMNS)


def test_every_extracted_field_has_a_value_column(sheet):
    names = header(sheet)
    for field in FIELD_NAMES:
        if field == "has_phd":
            continue
        assert export_name(field) in names


def test_no_evidence_columns_are_exported_only_the_two_course_names(sheet):
    names = header(sheet)
    assert not [n for n in names if n.endswith("_evidence")]
    assert evidence_column("highest_degree") == "highest_degree_course_name" in names
    assert evidence_column("phd_status") == "phd_course_name" in names


def test_has_phd_is_not_exported_but_phd_status_is(sheet):
    names = header(sheet)
    assert "has_phd" not in names and "has_phd_evidence" not in names
    assert "phd_status" in names and "phd_course_name" in names


def test_a_has_phd_review_reason_points_at_the_phd_status_column():
    from app.excel_writer import _format_reasons

    assert _format_reasons(["has_phd:low_confidence", "phd_status:low_confidence"]) == "phd_status (low confidence)"


def test_marks_column_is_spelled_out_and_marks_and_cgpa_have_no_evidence_column(sheet):
    names = header(sheet)
    assert "masters_percentage" in names and "masters_cgpa" in names
    assert not {"marks_pct", "marks_percentage", "cgpa"} & set(names)
    assert not [n for n in names if n.endswith("_evidence") and ("marks" in n or "cgpa" in n or "percentage" in n)]


def test_review_reasons_name_the_column_as_it_is_headed():
    from app.excel_writer import _format_reasons

    assert _format_reasons(["marks_pct:low_confidence"]) == "masters_percentage (low confidence)"
    assert _format_reasons(["cgpa:low_confidence"]) == "masters_cgpa (low confidence)"


def _marks_row(marks, cgpa):
    from llm.providers.fake_provider import CANNED_RESULTS

    result = CANNED_RESULTS[0].model_copy(deep=True)
    absent = {"value": None, "confidence": 0.0, "evidence": None}
    result.marks_pct = type(result.marks_pct)(**({"value": marks, "confidence": 0.9, "evidence": "x"} if marks is not None else absent))
    result.cgpa = type(result.cgpa)(**({"value": cgpa, "confidence": 0.9, "evidence": "x"} if cgpa is not None else absent))
    return build_row(ResumeRecord(source_filename="a.pdf", processed_at=datetime(2026, 10, 4), result=result))


def test_a_stated_percentage_is_shown_as_stated():
    row = _marks_row(63.56, None)
    assert row["masters_percentage"] == 63.56
    assert row["masters_percentage_source"] == "stated on resume"
    assert row["masters_cgpa"] is None


@pytest.mark.parametrize("cgpa, pct", [(8.2, 74.5), (9.13, 83.8), (7.78, 70.3), (8.79, 80.4), (10.0, 92.5)])
def test_a_cgpa_is_converted_by_the_aicte_formula_and_marked_as_converted(cgpa, pct):
    row = _marks_row(None, cgpa)
    assert row["masters_percentage"] == pytest.approx(pct)
    assert row["masters_percentage_source"] == "converted from CGPA"
    assert row["masters_cgpa"] == cgpa  # the stated figure is kept beside it


def test_a_stated_percentage_is_never_overwritten_by_a_conversion():
    row = _marks_row(76.6, 8.0)
    assert row["masters_percentage"] == 76.6
    assert row["masters_percentage_source"] == "stated on resume"


def test_no_marks_at_all_stays_blank_rather_than_inventing_a_percentage():
    row = _marks_row(None, None)
    assert row["masters_percentage"] is None and row["masters_percentage_source"] is None


def test_a_number_that_is_not_a_ten_point_cgpa_is_not_converted():
    from app.excel_writer import cgpa_to_percentage

    assert cgpa_to_percentage(63.5) is None
    assert cgpa_to_percentage(0) is None
    assert cgpa_to_percentage(0.5) == 0.0  # never negative


def test_no_confidence_columns_are_exported(sheet):
    """The number is internal: it drives needs_review and the highlight, but
    is not shown as a column anywhere."""
    assert not [n for n in header(sheet) if "confidence" in n.lower()]
    assert not [c for c in COLUMNS if "confidence" in c.lower()]


def test_rank_and_score_columns_exist_and_lead_the_sheet(sheet):
    """Ranking was added on an explicit instruction, overriding the original
    no-leaderboard design. See app/ranking.py for what the number is not."""
    names = header(sheet)
    assert names[:3] == ["rank", "score", "score_completeness"]


def test_rows_are_ordered_by_rank(sheet):
    ranks = [sheet.cell(row=r, column=1).value for r in range(2, sheet.max_row + 1)]
    present = [r for r in ranks if r is not None]
    assert present == sorted(present)
    assert present == list(range(1, len(present) + 1))


def test_score_completeness_travels_with_every_score(sheet):
    """A score without its completeness is the misleading form of this column:
    it hides that the number may rest on mostly-missing data."""
    names = header(sheet)
    si, ci = names.index("score") + 1, names.index("score_completeness") + 1
    for r in range(2, sheet.max_row + 1):
        if sheet.cell(row=r, column=si).value is not None:
            assert sheet.cell(row=r, column=ci).value is not None


def test_needs_review_still_present_alongside_rank(sheet):
    """A high rank on unverified data must stay visibly unverified."""
    assert "needs_review" in header(sheet)
    assert "review_reasons" in header(sheet)


def test_workbook_carries_the_ranking_disclaimer(records, tmp_path):
    """The caveat has to survive the file being forwarded to someone who
    never saw the app."""
    from openpyxl import load_workbook

    from app.ranking import RANKING_DISCLAIMER

    path = write_workbook(records, tmp_path / "d.xlsx", threshold=THRESHOLD)
    wb = load_workbook(path)
    assert "how_scoring_works" in wb.sheetnames
    explain = " ".join(
        str(c.value) for row in wb["how_scoring_works"].iter_rows() for c in row if c.value
    )
    assert "not a hiring decision" in explain
    note = wb[SHEET_NAME].cell(row=1, column=1).comment
    assert note is not None and "not a hiring decision" in note.text


def test_one_row_per_uploaded_file_including_the_failed_one(sheet, records):
    """Every uploaded file gets a row. Order is by rank now, not upload
    order, so compare as sets."""
    assert sheet.max_row == len(records) + 1
    names = header(sheet)
    fi = names.index("source_filename") + 1
    filenames = [sheet.cell(row=r, column=fi).value for r in range(2, sheet.max_row + 1)]
    assert sorted(filenames) == sorted(r.source_filename for r in records)


def test_header_row_is_frozen_and_bold_with_autofilter(sheet):
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None
    assert sheet.auto_filter.ref.startswith("A1:")
    assert sheet.auto_filter.ref.endswith(str(sheet.max_row))
    assert all(cell.font.bold for cell in sheet[1])


def test_column_widths_are_set_not_left_at_default(sheet):
    widths = sheet.column_dimensions
    for i in range(1, len(COLUMNS) + 1):
        letter = widths[list(widths)[0]].__class__  # noqa: F841 - keep openpyxl import local
    from openpyxl.utils import get_column_letter

    # "rank" is a 1-2 digit column; a wide one would just be padding.
    narrow = {"rank", "score"}
    for i in range(1, len(COLUMNS) + 1):
        dim = sheet.column_dimensions[get_column_letter(i)]
        floor = 5 if COLUMNS[i - 1] in narrow else 8
        assert dim.width and dim.width > floor


def test_low_confidence_values_are_highlighted_in_place(sheet, records):
    """With no confidence column, the doubtful value itself carries the fill.
    Blank cells are never filled: absent is not the same as doubtful."""
    by_name = {r.source_filename: r for r in records}
    name_col = column_index(sheet, "source_filename")
    filled, unfilled = 0, 0
    for field in FIELD_NAMES:
        if field in ("candidate_name", "has_phd"):
            continue
        col = column_index(sheet, export_name(field))
        for r in range(2, sheet.max_row + 1):
            record = by_name[sheet.cell(row=r, column=name_col).value]
            cell = sheet.cell(row=r, column=col)
            has_fill = cell.fill is not None and cell.fill.start_color.rgb == "00FFF2CC"
            if cell.value is None or record.result is None:
                assert not has_fill, f"{field} row {r} is blank but highlighted"
                continue
            # A converted percentage is only as doubtful as the CGPA behind it.
            converted = field == "marks_pct" and record.result.marks_pct.value is None
            if getattr(record.result, "cgpa" if converted else field).confidence < THRESHOLD:
                assert has_fill, f"{field} row {r} below threshold but not highlighted"
                filled += 1
            else:
                assert not has_fill, f"{field} row {r} above threshold but highlighted"
                unfilled += 1
    assert filled > 0 and unfilled > 0


def test_needs_review_is_a_real_boolean_not_a_string(sheet):
    col = column_index(sheet, "needs_review")
    values = [sheet.cell(row=r, column=col).value for r in range(2, sheet.max_row + 1)]
    assert all(isinstance(v, bool) for v in values)
    assert True in values and False in values


def test_failed_file_row_is_empty_with_a_parse_error_and_needs_review(sheet):
    last = sheet.max_row
    row = {name: sheet.cell(row=last, column=i + 1).value for i, name in enumerate(COLUMNS)}
    assert row["source_filename"] == "scanned_image.pdf"
    assert row["needs_review"] is True
    assert row["parse_error"] == "no extractable text in file"
    assert row["candidate_name"] is None
    assert all(row[c] is None for c in EVIDENCE_COLUMNS)


def test_dates_are_written_as_dates_so_the_column_sorts_chronologically(sheet):
    col = column_index(sheet, "masters_award_date")
    values = [sheet.cell(row=r, column=col).value for r in range(2, sheet.max_row + 1)]
    written = [v for v in values if v is not None]
    assert written, "expected at least one date from the canned results"
    assert all(isinstance(v, (date, datetime)) for v in written)


def test_absent_field_writes_blank_value_and_no_evidence():
    """Absent means empty — never a plausible-looking default."""
    record = ResumeRecord(
        source_filename="a.pdf",
        processed_at=datetime(2026, 8, 22),
        result=FakeProvider().extract_fields("x"),
    )
    row = build_row(record, threshold=THRESHOLD)
    assert row["phd_award_date"] is None
    assert "phd_award_date_confidence" not in row
    assert "phd_award_date_evidence" not in row


def test_set_without_state_row_is_flagged_for_review(sheet):
    """The second canned result claims SET with no grounded state."""
    col = column_index(sheet, "needs_review")
    status_col = column_index(sheet, "net_set_status")
    for r in range(2, sheet.max_row + 1):
        if sheet.cell(row=r, column=status_col).value in ("SET", "SLET"):
            state = sheet.cell(row=r, column=column_index(sheet, "set_state")).value
            if state is None:
                assert sheet.cell(row=r, column=col).value is True


def test_the_quotes_are_still_available_in_the_raw_output_sheet(records, tmp_path):
    """Dropping the evidence columns must not drop the evidence: the model's
    own response, quotes included, stays on the audit sheet."""
    path = write_workbook(records, tmp_path / "out.xlsx", threshold=THRESHOLD)
    raw = load_workbook(path)["raw_llm_output"]
    cells = [raw.cell(row=r, column=2).value or "" for r in range(2, raw.max_row + 1)]
    assert any("UGC-NET" in c for c in cells)


def test_review_reasons_column_exists_and_follows_needs_review():
    """needs_review says *that* a row needs checking; review_reasons says
    *which fields*, so a reviewer isn't hunting across 30+ columns."""
    from app.excel_writer import COLUMNS

    assert "review_reasons" in COLUMNS
    assert COLUMNS.index("review_reasons") == COLUMNS.index("needs_review") + 1


def test_review_reasons_names_the_failing_fields():
    from datetime import datetime

    from app.excel_writer import build_row
    from llm.interface import ResumeRecord
    from tests.test_confidence_routing import _ok, build

    result = build(publications_count=_ok(3, 0.95, None))  # claimed, unquoted
    record = ResumeRecord(
        source_filename="cv.pdf", processed_at=datetime.now(), result=result
    )
    row = build_row(record)
    assert row["needs_review"] is True
    assert "publications_count" in row["review_reasons"]
    assert "no evidence" in row["review_reasons"]


def test_review_reasons_is_blank_for_a_clean_row():
    from datetime import datetime

    from app.excel_writer import build_row
    from llm.interface import ResumeRecord
    from tests.test_confidence_routing import build

    record = ResumeRecord(
        source_filename="cv.pdf", processed_at=datetime.now(), result=build()
    )
    row = build_row(record)
    assert row["needs_review"] is False
    assert row["review_reasons"] is None


def test_review_reasons_carries_no_extracted_values():
    """This column gets shared with reviewers; it must name fields and rules,
    never the candidate's data."""
    from datetime import datetime

    from app.excel_writer import build_row
    from llm.interface import ResumeRecord
    from tests.test_confidence_routing import _ok, build

    result = build(
        net_set_status=_ok("SET"), set_state=_ok("Maharashtra", 0.2)
    )
    record = ResumeRecord(
        source_filename="cv.pdf", processed_at=datetime.now(), result=result
    )
    row = build_row(record)
    assert "Maharashtra" not in (row["review_reasons"] or "")


# --- date precision is disclosed, not fabricated (P2) -----------------------

def test_a_year_only_date_is_not_written_as_a_full_iso_date(tmp_path):
    """A bare "2015" parses to date(2015,1,1) and then reads in the export
    exactly like a day-accurate date. The output must not assert a January
    1st the resume never mentioned."""
    from datetime import datetime as _dt

    from app.excel_writer import build_row
    from tests.test_confidence_routing import _ok, build

    result = build(masters_award_date=_ok(date(2015, 1, 1), 0.85, "M.Tech 2015"))
    result.masters_award_date_precision = "year"
    row = build_row(
        ResumeRecord(source_filename="cv.pdf", processed_at=_dt.now(), result=result)
    )
    assert row["masters_award_date"] == "2015"


def test_a_month_precision_date_shows_year_and_month_only(tmp_path):
    from datetime import datetime as _dt

    from app.excel_writer import build_row
    from tests.test_confidence_routing import _ok, build

    result = build(masters_award_date=_ok(date(2021, 5, 1), 0.9, "Graduated: May 2021"))
    result.masters_award_date_precision = "month"
    row = build_row(
        ResumeRecord(source_filename="cv.pdf", processed_at=_dt.now(), result=result)
    )
    assert row["masters_award_date"] == "05-2021"


def test_a_genuinely_full_date_is_still_written_as_a_date(tmp_path):
    from datetime import datetime as _dt

    from app.excel_writer import build_row
    from tests.test_confidence_routing import _ok, build

    result = build(masters_award_date=_ok(date(2019, 11, 12), 0.9, "awarded 12 November 2019"))
    result.masters_award_date_precision = "full"
    row = build_row(
        ResumeRecord(source_filename="cv.pdf", processed_at=_dt.now(), result=result)
    )
    assert row["masters_award_date"] == date(2019, 11, 12)


# --- raw_llm_output is retrievable for audit (P2) ---------------------------

def test_raw_llm_output_is_retrievable_from_the_workbook(records, tmp_path):
    """The parsed columns are the product of a prompt, a schema and several
    post-processing rules. Only the model's own response distinguishes "the
    model said something false" from "we mangled something true"."""
    from openpyxl import load_workbook

    path = write_workbook(records, tmp_path / "out.xlsx", threshold=THRESHOLD)
    wb = load_workbook(path)
    assert "raw_llm_output" in wb.sheetnames
    sheet = wb["raw_llm_output"]
    filenames = [c.value for c in sheet["A"][1:]]
    for rec in records:
        assert rec.source_filename in filenames


def test_a_failed_extraction_still_gets_a_raw_output_row(tmp_path):
    from datetime import datetime as _dt

    from openpyxl import load_workbook

    recs = [ResumeRecord.failed("scanned.pdf", "no extractable text")]
    path = write_workbook(recs, tmp_path / "out.xlsx", threshold=THRESHOLD)
    sheet = load_workbook(path)["raw_llm_output"]
    assert sheet["A2"].value == "scanned.pdf"
    assert sheet["B2"].value  # carries the reason rather than being blank


# --- one date convention: DD-MM-YYYY ----------------------------------------


def test_full_dates_are_real_dates_shown_day_month_year(sheet):
    col = column_index(sheet, "masters_award_date")
    cells = [sheet.cell(row=r, column=col) for r in range(2, sheet.max_row + 1)]
    dated = [c for c in cells if isinstance(c.value, (date, datetime))]
    assert dated and all(c.number_format == "dd-mm-yyyy" for c in dated)


def test_processed_at_uses_the_same_convention(sheet):
    col = column_index(sheet, "processed_at")
    assert all(
        sheet.cell(row=r, column=col).number_format == "dd-mm-yyyy hh:mm" for r in range(2, sheet.max_row + 1)
    )


def test_the_dashboard_table_writes_dates_in_the_same_convention():
    from app.excel_writer import date_text, rows_for_display

    assert date_text(date(2019, 11, 12)) == "12-11-2019"
    assert date_text(datetime(2026, 10, 4, 12, 54, 30)) == "04-10-2026 12:54"
    assert date_text("05-2021") == "05-2021" and date_text("2015") == "2015" and date_text(None) is None
    shown = rows_for_display([{"phd_award_date": date(2019, 11, 12), "score": 61.2, "needs_review": True}])
    assert shown == [{"phd_award_date": "12-11-2019", "score": 61.2, "needs_review": True}]


def test_accuracy_scoring_reads_the_new_convention_against_iso_answer_keys():
    from tools.score_accuracy import values_match

    assert values_match("2019-11-12", "12-11-2019")
    assert values_match("2021-05-01", "05-2021")  # key padded the day; export states the month
    assert values_match("2015-01-01", "2015")
    assert not values_match("2019-11-12", "11-12-2019")  # day and month are not interchangeable
    assert not values_match("2021-05-01", "06-2021")
