"""Tests for the batch runner: ordering, concurrency, and failure isolation.

No Streamlit and no real model — process_batch takes a provider and plain
bytes, which is exactly why it was decoupled from the upload object.
"""

from __future__ import annotations

import threading
import time

import pytest
from docx import Document

from app.dashboard import process_batch
from llm.interface import ExtractionFailure
from llm.providers.fake_provider import CANNED_RESULTS


def _docx_bytes(name: str) -> bytes:
    import io

    d = Document()
    d.add_paragraph(f"{name} Synthetic Candidate")
    d.add_paragraph("M.Sc. Physics, Synthetic University, 2015 (68.4%)")
    d.add_paragraph("UGC-NET (Physical Sciences), December 2016")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class _RecordingProvider:
    """Returns a canned result, recording peak concurrent in-flight calls."""

    name = "recording"

    def __init__(self, max_concurrency: int, delay: float = 0.05) -> None:
        self.max_concurrency = max_concurrency
        self._delay = delay
        self._lock = threading.Lock()
        self._in_flight = 0
        self.peak_in_flight = 0
        self.seen: list[str] = []

    def extract_fields(self, resume_text: str):
        with self._lock:
            self._in_flight += 1
            self.peak_in_flight = max(self.peak_in_flight, self._in_flight)
            self.seen.append(resume_text[:20])
        try:
            time.sleep(self._delay)
            return CANNED_RESULTS[0].model_copy(deep=True)
        finally:
            with self._lock:
                self._in_flight -= 1


class _ExplodingProvider:
    """Fails on one specific resume, succeeds on the rest."""

    name = "exploding"
    max_concurrency = 4

    def __init__(self, fail_marker: str) -> None:
        self._fail_marker = fail_marker

    def extract_fields(self, resume_text: str):
        if self._fail_marker in resume_text:
            raise ExtractionFailure("simulated extraction failure")
        return CANNED_RESULTS[0].model_copy(deep=True)


UPLOADS = [(f"cv{i}.docx", _docx_bytes(f"Cand{i}")) for i in range(6)]


def test_results_are_returned_in_upload_order_not_completion_order():
    """Rows must line up with the order files were uploaded, even though
    workers finish out of order."""
    provider = _RecordingProvider(max_concurrency=4)
    records = process_batch(provider, UPLOADS)
    assert [r.source_filename for r in records] == [name for name, _ in UPLOADS]


def test_every_upload_produces_exactly_one_record():
    provider = _RecordingProvider(max_concurrency=4)
    records = process_batch(provider, UPLOADS)
    assert len(records) == len(UPLOADS)


def test_concurrency_is_capped_at_provider_limit():
    provider = _RecordingProvider(max_concurrency=3)
    process_batch(provider, UPLOADS)
    assert provider.peak_in_flight <= 3


def test_serial_provider_never_overlaps_calls():
    """A local CPU model declares max_concurrency=1; overlapping calls there
    would thrash the same cores."""
    provider = _RecordingProvider(max_concurrency=1)
    process_batch(provider, UPLOADS)
    assert provider.peak_in_flight == 1


def test_parallel_provider_actually_overlaps():
    provider = _RecordingProvider(max_concurrency=4, delay=0.1)
    process_batch(provider, UPLOADS)
    assert provider.peak_in_flight > 1


def test_one_failing_resume_does_not_sink_the_batch():
    provider = _ExplodingProvider(fail_marker="Cand2")
    records = process_batch(provider, UPLOADS)
    assert len(records) == len(UPLOADS)
    failed = [r for r in records if r.parse_error]
    assert len(failed) == 1
    assert failed[0].source_filename == "cv2.docx"
    # every other row still extracted normally
    assert all(r.result is not None for r in records if not r.parse_error)


def test_progress_callback_reports_each_completion():
    provider = _RecordingProvider(max_concurrency=4)
    seen: list[tuple[int, int]] = []
    process_batch(provider, UPLOADS, on_progress=lambda d, t: seen.append((d, t)))
    assert len(seen) == len(UPLOADS)
    assert seen[-1] == (len(UPLOADS), len(UPLOADS))
    assert [d for d, _ in seen] == sorted(d for d, _ in seen)


def test_unreadable_file_still_produces_a_row():
    provider = _RecordingProvider(max_concurrency=4)
    uploads = [("broken.pdf", b"not a real pdf at all")]
    records = process_batch(provider, uploads)
    assert len(records) == 1
    assert records[0].parse_error
    assert records[0].result is None


@pytest.mark.parametrize("limit", [1, 2, 8])
def test_worker_count_never_exceeds_upload_count(limit):
    provider = _RecordingProvider(max_concurrency=limit)
    records = process_batch(provider, UPLOADS[:2])
    assert len(records) == 2
    assert provider.peak_in_flight <= 2
