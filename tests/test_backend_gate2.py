"""Phase 6: Gate 2 (approve, override, send back), the outcome list and the Excel download.

Fake provider, synthetic data, no network. SQLite by default; PostgreSQL when TEST_DATABASE_URL is set.
"""

from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from backend import gate1, gate2, states
from backend.assessor_service import assess_application, assess_opening, latest_evaluation
from backend.export import CANDIDATE_COLUMNS, opening_workbook
from backend.models import Application, CandidateQualification, EvaluationResult, HrDecision, ReviewEdit
from backend.reader_service import read_application
from backend.rules_seed import seed_rules
from backend.web import STATE_LABELS
from llm.interface import EducationEntry
from llm.providers.fake_provider import CANNED_RESULTS, FakeProvider
from tests.test_backend_foundation import _application, _clean_result, client, engine, session as _plain_session  # noqa: F401
from tests.test_backend_intake import _opening, _storage  # noqa: F401

TODAY = date(2026, 10, 6)
_people = iter(range(1, 10_000))


@pytest.fixture
def session(_plain_session):
    seed_rules(_plain_session)
    _plain_session.commit()
    return _plain_session


def _assessed(session, tmp_path, result=None, **opening_fields) -> Application:
    """An application read cleanly and assessed. Default: UGC opening, M.Sc. 68.4% with NET -> SHORTLISTED."""
    o = _opening(session, **{"school_id": "SCH-013", "discipline_group": "GENERAL", **opening_fields})
    a = _application(session, tmp_path, category="General", differently_abled=False, study_leave_taken=False)
    a.opening_id, a.school_id, a.applied_designation = o.opening_id, o.school_id, o.designation
    read_application(session, a, FakeProvider(script=[result or _clean_result()]))
    assert a.status == states.EXTRACTED
    assess_application(session, a, TODAY)
    session.commit()
    return a


def _no_net():
    r = _clean_result()
    r.net_set_status.value = "NONE"
    return r


def _other_person():
    """The clean result under another made-up address, so two applications are not taken for one person."""
    r = _clean_result()
    r.email = f"person{next(_people)}@example.org"
    return r


def _class_unstated():
    r = _clean_result()
    r.marks_pct.value = r.marks_pct.evidence = None
    r.education = [EducationEntry(level="UG", degree="B.E."), EducationEntry(level="PG", degree="M.E.")]
    return r


# --- approve -----------------------------------------------------------------


def test_approving_records_the_decision_against_the_assessment_and_moves_on(session, tmp_path):
    a = _assessed(session, tmp_path)
    d = gate2.approve(session, a)
    assert a.status == states.HR_APPROVED
    assert (d.action, d.final_outcome, d.final_designation, d.justification) == ("APPROVED", "SHORTLISTED", "ASSISTANT_PROFESSOR", None)
    assert d.evaluation_id == latest_evaluation(session, a.application_id).evaluation_id
    last = a.transitions[-1]
    assert (last.from_state, last.to_state, last.actor, last.note) == ("SHORTLISTED", "HR_APPROVED", "user:hr", "gate2: approved as assessed")
    assert gate2.final_decision(session, a) is d


def test_approving_a_not_eligible_finding_records_it_as_such(session, tmp_path):
    a = _assessed(session, tmp_path, _no_net())
    assert a.status == states.NOT_ELIGIBLE
    d = gate2.approve(session, a)
    assert (d.final_outcome, d.final_designation) == ("NOT_ELIGIBLE", None)


def test_there_is_nothing_to_approve_when_the_rules_could_not_settle_it(session, tmp_path):
    a = _assessed(session, tmp_path, _class_unstated(), school_id="SCH-008", discipline_group="ENGINEERING_TECHNOLOGY")
    assert a.status == states.MANUAL_REVIEW
    with pytest.raises(gate2.DecisionError, match="no finding to approve"):
        gate2.approve(session, a)
    d = gate2.override(session, a, "SHORTLISTED", "ASSISTANT_PROFESSOR", "Marksheets seen: First Class in the M.E.")
    assert d.action == "DECIDED" and a.status == states.HR_APPROVED and a.transitions[-1].note == "gate2: decided by HR"


# --- override ----------------------------------------------------------------


def test_an_override_needs_a_decision_a_post_and_a_reason(session, tmp_path):
    a = _assessed(session, tmp_path, _no_net())
    for args, field in (
        (("", None, "Certificate of NET seen in original."), "outcome"),
        (("SHORTLISTED", None, "Certificate of NET seen in original."), "designation"),
        (("SHORTLISTED", "DEAN", "Certificate of NET seen in original."), "designation"),
        (("SHORTLISTED", "ASSISTANT_PROFESSOR", "ok"), "justification"),
        (("SHORTLISTED", "ASSISTANT_PROFESSOR", "x" * 1001), "justification"),
    ):
        with pytest.raises(gate2.DecisionError) as exc:
            gate2.override(session, a, *args)
        assert field in exc.value.errors
    assert a.status == states.NOT_ELIGIBLE and session.scalars(select(HrDecision)).all() == []


def test_an_override_is_kept_beside_the_finding_it_overrode_and_its_reason_stays_out_of_the_audit_trail(session, tmp_path):
    a = _assessed(session, tmp_path, _no_net())
    reason = "NET certificate (December 2016) seen in original; the resume omitted it."
    d = gate2.override(session, a, "SHORTLISTED", "ASSISTANT_PROFESSOR", f"  {reason}  ")
    assert (d.action, d.final_outcome, d.final_designation, d.justification) == ("OVERRIDDEN", "SHORTLISTED", "ASSISTANT_PROFESSOR", reason)
    assert a.status == states.HR_APPROVED
    assert latest_evaluation(session, a.application_id).outcome == "NOT_ELIGIBLE"  # the engine's finding is untouched
    assert a.transitions[-1].note == "gate2: overridden with a recorded justification"
    assert all("certificate" not in (t.note or "") for t in a.transitions)


def test_overriding_to_not_eligible_records_no_post(session, tmp_path):
    a = _assessed(session, tmp_path)
    d = gate2.override(session, a, "NOT_ELIGIBLE", "PROFESSOR", "Degree certificate could not be verified with the university.")
    assert (d.final_outcome, d.final_designation) == ("NOT_ELIGIBLE", None)


def test_a_decision_cannot_be_made_twice_or_before_assessment(session, tmp_path):
    a = _assessed(session, tmp_path)
    gate2.approve(session, a)
    for attempt in (lambda: gate2.approve(session, a),
                    lambda: gate2.override(session, a, "NOT_ELIGIBLE", None, "Changed my mind about this one."),
                    lambda: gate2.return_for_reassessment(session, a)):
        with pytest.raises(gate2.DecisionError, match="not waiting for an HR decision"):
            attempt()
    assert len(session.scalars(select(HrDecision)).all()) == 1


# --- send back ---------------------------------------------------------------


def test_sent_back_for_correction_the_named_fields_reopen_and_the_next_assessment_uses_them(session, tmp_path):
    a = _assessed(session, tmp_path, _no_net())
    with pytest.raises(gate2.DecisionError):
        gate2.return_for_correction(session, a, ["not_a_field"])
    gate2.return_for_correction(session, a, ["net_set_status"])
    assert a.status == states.PENDING_REVIEW and a.transitions[-1].note == "gate2: returned for net_set_status"
    flags = gate1.flagged_fields(session, a)
    assert [f.name for f in flags] == ["net_set_status", "set_state"]
    assert flags[0].reasons == ["Returned by HR for correction after the assessment."]

    gate1.save_review(session, a, {"net_set_status": "NET"}, set())
    assert a.status == states.EXTRACTED
    assert assess_opening(session, a.opening_id, TODAY) == {"SHORTLISTED": 1}
    evaluations = session.scalars(select(EvaluationResult).order_by(EvaluationResult.evaluation_id)).all()
    assert [e.outcome for e in evaluations] == ["NOT_ELIGIBLE", "SHORTLISTED"]  # the first is kept, not rewritten
    assert latest_evaluation(session, a.application_id).outcome == "SHORTLISTED"
    assert [d.action for d in gate2.decisions(session, a.application_id)] == ["RETURNED"]
    assert gate2.final_decision(session, a) is None
    gate2.approve(session, a)
    assert gate2.final_decision(session, a).action == "APPROVED"


def test_the_bachelors_marks_can_be_entered_when_the_class_was_not_on_the_resume(session, tmp_path):
    a = _assessed(session, tmp_path, _class_unstated(), school_id="SCH-008", discipline_group="ENGINEERING_TECHNOLOGY")
    assert a.status == states.MANUAL_REVIEW
    gate2.return_for_correction(session, a, ["ug_marks_pct"])
    assert gate1.flagged_fields(session, a)[0].spec.label == "Bachelor's marks (percentage)"
    gate1.save_review(session, a, {"ug_marks_pct": "64.5"}, set())
    ug = session.scalars(select(CandidateQualification).where(CandidateQualification.degree_level == "UG")).one()
    assert ug.marks_pct == 64.5
    assert session.scalars(select(ReviewEdit).where(ReviewEdit.field == "ug_marks_pct")).one().action == "ENTERED"
    assert assess_application(session, a, TODAY).outcome == "SHORTLISTED"  # First Class in one degree (AICTE cl. 5.1(a))


def test_looking_at_the_form_does_not_add_a_degree_row_and_saving_does(session, tmp_path):
    r = _clean_result()
    r.education = [EducationEntry(level="PG", degree="M.Sc.", marks_pct=68.4)]
    a = _assessed(session, tmp_path, r)
    gate2.return_for_correction(session, a, ["ug_cgpa"])
    assert gate1.flagged_fields(session, a)[0].value == ""
    session.flush()
    assert [q.degree_level for q in session.scalars(select(CandidateQualification))] == ["PG"]
    gate1.save_review(session, a, {"ug_cgpa": "7.9"}, set())
    added = session.scalars(select(CandidateQualification).where(CandidateQualification.degree_level == "UG")).one()
    assert added.cgpa == 7.9 and added.found_in_resume is False  # entered by a person, and marked so


def test_sent_back_for_reassessment_as_it_stands(session, tmp_path):
    a = _assessed(session, tmp_path)
    gate2.return_for_reassessment(session, a)
    assert a.status == states.EXTRACTED and a.transitions[-1].note == "gate2: returned for re-assessment"
    assert assess_application(session, a, TODAY).outcome == "SHORTLISTED"
    assert len(session.scalars(select(EvaluationResult)).all()) == 2


# --- moving an application to another opening --------------------------------


def test_moving_an_assessed_application_takes_the_new_openings_post_and_rules_and_is_assessed_again(session, tmp_path):
    from backend import intake

    a = _assessed(session, tmp_path, None, designation="PROFESSOR")
    assert a.status == states.RE_CATEGORISED  # filed as a Professor applicant: no Ph.D., so only a lower post is met
    target = _opening(session, school_id="SCH-008", discipline_group="ENGINEERING_TECHNOLOGY", designation="ASSISTANT_PROFESSOR")
    intake.move_application(session, a, target)
    assert (a.opening_id, a.school_id, a.applied_designation, a.status) == (target.opening_id, "SCH-008", "ASSISTANT_PROFESSOR", states.EXTRACTED)
    assert a.transitions[-1].note == f"moved to {target.reference}; to be assessed again"
    moved = session.scalars(select(ReviewEdit).where(ReviewEdit.field == "opening")).one()
    assert (moved.action, moved.new_value) == ("MOVED", target.reference)
    assert [d.action for d in gate2.decisions(session, a.application_id)] == ["RETURNED"]
    d = assess_application(session, a, TODAY)
    assert d.outcome == "SHORTLISTED" and d.rule_version == "AICTE-DEGREE-2019"  # First Class in both degrees
    assert [e.outcome for e in session.scalars(select(EvaluationResult).order_by(EvaluationResult.evaluation_id))] == ["RE_CATEGORISED", "SHORTLISTED"]


def test_an_application_cannot_be_moved_once_decided_or_onto_itself_or_a_closed_opening_or_its_own_twin(session, tmp_path):
    from backend import intake

    a = _assessed(session, tmp_path)
    here, there = a.opening, _opening(session, school_id="SCH-008")
    with pytest.raises(intake.IntakeError, match="already under"):
        intake.move_application(session, a, here)
    twin = _application(session, tmp_path, category="General")  # the same resume file, already under the target
    twin.opening_id = there.opening_id
    session.flush()
    with pytest.raises(intake.IntakeError, match="already has this candidate or this resume"):
        intake.move_application(session, a, there)
    twin.status = states.WITHDRAWN
    session.flush()
    intake.close_opening(session, there)
    with pytest.raises(intake.IntakeError, match="closed"):
        intake.move_application(session, a, there)
    gate2.approve(session, a)
    with pytest.raises(intake.IntakeError, match="cannot be moved"):
        intake.move_application(session, a, _opening(session, school_id="SCH-009"))


def test_moving_a_whole_opening_moves_what_it_can_and_says_why_the_rest_stayed(session, tmp_path):
    from backend import intake

    decided = _assessed(session, tmp_path)
    source = decided.opening
    gate2.approve(session, decided)
    waiting = _application(session, tmp_path, data=b"%PDF-1.4 another made-up file", filename="other.pdf")
    waiting.opening_id = source.opening_id
    session.flush()
    target = _opening(session, school_id="SCH-008")
    moved, stayed = intake.move_all(session, source, target)
    assert moved == 1 and waiting.opening_id == target.opening_id and waiting.status == states.RECEIVED
    assert stayed == [f"{decided.reference}: An application that is HR_APPROVED cannot be moved."]


def test_moving_from_the_pages(client, engine, tmp_path):
    from sqlalchemy.orm import Session as _Session

    app_id, opening_id = _assessed_on_server(engine, tmp_path)
    with _Session(engine) as s:
        target = _opening(s, school_id="SCH-008", discipline_group="ENGINEERING_TECHNOLOGY").opening_id
    page = client.get(f"/hr/applications/{app_id}").text
    assert "Filed under the wrong post?" in page and f'value="{target}"' in page
    moved = client.post(f"/hr/applications/{app_id}/move", data={"opening_id": target}).text
    assert f"Moved to OPN-{target:05d}" in moved and "An earlier assessment was sent back" in moved
    assert "Choose an opening" in client.post(f"/hr/applications/{app_id}/move", data={"opening_id": 999999}).text
    back = client.post(f"/hr/openings/{target}/move-all", data={"target_id": opening_id}).text
    assert f"Moved 1 application(s) to OPN-{opening_id:05d}." in back
    assert "Choose another opening" in client.post(f"/hr/openings/{target}/move-all", data={"target_id": target}).text


# --- the Reader no longer asks for what the opening's rules never use --------


def test_an_aicte_opening_is_not_asked_for_the_phd_regulations_or_a_set_state(session, tmp_path):
    def read(group, school):
        o = _opening(session, school_id=school, discipline_group=group)
        a = _application(session, tmp_path, category="General", differently_abled=False, study_leave_taken=False)
        a.opening_id = o.opening_id
        read_application(session, a, FakeProvider(script=[CANNED_RESULTS[1].model_copy(deep=True)]))  # Ph.D., SET with no State
        return set(a.extracted.review_reasons)

    ugc, aicte = read("GENERAL", "SCH-013"), read("ENGINEERING_TECHNOLOGY", "SCH-008")
    asked_only_under_ugc = {"phd_regulation:manual_entry_required", "set_state:required_for_set"}
    assert asked_only_under_ugc <= ugc and not (asked_only_under_ugc & aicte)
    assert aicte == ugc - asked_only_under_ugc


# --- the Excel download ------------------------------------------------------


def _sheet(wb, name) -> list[dict]:
    rows = list(wb[name].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def test_the_workbook_has_the_sheets_hr_knows_keyed_by_reference_in_order_of_receipt(session, tmp_path):
    first = _assessed(session, tmp_path, _no_net())
    second = _application(session, tmp_path, category="SC", differently_abled=False, study_leave_taken=False)
    second.opening_id = first.opening_id
    other_person = _clean_result()
    other_person.email = "another.exampleton@example.org"
    read_application(session, second, FakeProvider(script=[other_person]))
    assess_application(session, second, TODAY)
    gate2.override(session, first, "SHORTLISTED", "ASSISTANT_PROFESSOR", "NET certificate seen in original.")
    session.commit()

    wb = load_workbook(io.BytesIO(opening_workbook(session, first.opening, STATE_LABELS)))
    assert wb.sheetnames == ["candidates", "education", "publications", "seminars_workshops", "teaching_skills", "experience",
                             "patents_awards_projects", "guidance_memberships", "contact_details", "assessment_checks", "how_to_read"]
    candidates = _sheet(wb, "candidates")
    assert tuple(candidates[0]) == CANDIDATE_COLUMNS and "rank" not in CANDIDATE_COLUMNS and "score" not in CANDIDATE_COLUMNS
    assert [c["reference"] for c in candidates] == [first.reference, second.reference]
    one, two = candidates
    assert one["finding"] == "Does not meet the minimum qualifications" and "cl. 3.3" in one["clause"] and one["page"]
    assert (one["hr_decision"], one["decided_post"], one["hr_justification"]) == ("Overridden", "Assistant Professor", "NET certificate seen in original.")
    assert two["finding"] == "Meets the post applied for" and two["hr_decision"] is None and two["category"] == "SC"
    assert (two["highest_degree"], two["highest_degree_course_name"]) == ("PG", "M.Sc. Physics")
    assert (two["masters_percentage"], two["masters_percentage_source"]) == (68.4, "stated on resume")
    assert two["shortlisting_score_table_3a"] and two["phd_status"] == "NA" and two["study_leave_taken"] == "No"

    education = _sheet(wb, "education")
    assert [(r["reference"], r["level"], r["degree_and_course"]) for r in education if r["reference"] == second.reference] == [
        (second.reference, "Bachelor's", "B.Sc. Physics"), (second.reference, "Master's", "M.Sc. Physics")]
    assert _sheet(wb, "experience")[0]["to"] == "Present" and _sheet(wb, "experience")[0]["type"] == "Teaching"
    assert _sheet(wb, "publications")[0]["title"] == "Paper A" and _sheet(wb, "teaching_skills")[0]["skills"] == "Python"
    assert {r["email"] for r in _sheet(wb, "contact_details")} == {"sample.exampleton@example.org", "another.exampleton@example.org"}
    checks = _sheet(wb, "assessment_checks")
    assert "Master's marks" in {c["requirement"] for c in checks} and {c["result"] for c in checks} <= {"Met", "Not met", "Open"}
    assert "Nothing here is a ranking" in " ".join(str(r[0]) for r in wb["how_to_read"].iter_rows(values_only=True))
    for name in wb.sheetnames[:-1]:  # the layout HR is used to: grey bold header, frozen key columns, filters
        ws = wb[name]
        assert ws.freeze_panes == "C2" and ws["A1"].font.b and ws["A1"].fill.fgColor.rgb.endswith("E8E8E8") and ws.auto_filter.ref


def test_a_cgpa_is_shown_as_a_percentage_for_reading_and_marked_as_converted(session, tmp_path):
    r = _clean_result()
    r.marks_pct.value = r.marks_pct.evidence = None
    r.cgpa.value, r.cgpa.confidence, r.cgpa.evidence = 8.2, 0.9, "CGPA 8.2"
    a = _assessed(session, tmp_path, r)
    row = _sheet(load_workbook(io.BytesIO(opening_workbook(session, a.opening, STATE_LABELS))), "candidates")[0]
    assert (row["masters_percentage"], row["masters_percentage_source"], row["masters_cgpa"]) == (74.5, "converted from CGPA", 8.2)
    assert row["finding"] == "For a person to decide"  # the converted figure decided nothing


def test_rows_waiting_for_a_person_are_shaded_and_a_returned_assessment_is_not_reported(session, tmp_path):
    a = _assessed(session, tmp_path, _no_net())
    gate2.return_for_correction(session, a, ["net_set_status"])
    session.commit()
    wb = load_workbook(io.BytesIO(opening_workbook(session, a.opening, STATE_LABELS)))
    row = _sheet(wb, "candidates")[0]
    assert row["finding"] is None and row["fields_to_check"] == "net_set_status (returned by hr)"
    assert wb["candidates"]["A2"].fill.fgColor.rgb.endswith("FFF2CC") and _sheet(wb, "assessment_checks") == []


def test_text_from_a_resume_cannot_become_a_spreadsheet_formula(session, tmp_path):
    a = _assessed(session, tmp_path)
    a.applicant_name = "=HYPERLINK(\"http://example.org\",\"click\")"
    session.commit()
    sheet = load_workbook(io.BytesIO(opening_workbook(session, a.opening, STATE_LABELS)))["candidates"]
    assert sheet["B2"].value.startswith("'=") and sheet["B2"].data_type == "s"


# --- the pages ---------------------------------------------------------------


def _assessed_on_server(engine, tmp_path, result=None):
    from sqlalchemy.orm import Session as _Session

    with _Session(engine) as s:
        seed_rules(s)
        a = _assessed(s, tmp_path, result)
        return a.application_id, a.opening_id


def test_the_application_page_offers_the_three_choices_and_records_an_approval(client, engine, tmp_path):
    app_id, opening_id = _assessed_on_server(engine, tmp_path)
    page = client.get(f"/hr/applications/{app_id}").text
    assert "Your decision" in page and "Approve as assessed" in page and "Or override it" in page and "Or send it back" in page
    done = client.post(f"/hr/applications/{app_id}/approve").text
    assert "Approved as assessed." in done and "HR decision" in done and "Your decision" not in done
    assert "Selection is made by the Selection Committee" in done
    listing = client.get(f"/hr/openings/{opening_id}").text
    assert "Approved:" in listing and "short-listed, Assistant Professor" in listing and "Decided by HR" in listing
    assert "not ranked" in listing


def test_an_override_without_a_reason_is_refused_from_the_page_and_one_with_it_is_shown(client, engine, tmp_path):
    app_id, opening_id = _assessed_on_server(engine, tmp_path, _no_net())
    listing = client.get(f"/hr/openings/{opening_id}").text
    assert "no NET" in listing and "cl. 3.3" in listing and ">Decide</a>" in listing
    refused = client.post(f"/hr/applications/{app_id}/override", data={"outcome": "SHORTLISTED", "designation": "ASSISTANT_PROFESSOR"}).text
    assert "Give the reason in at least 15 characters" in refused and "Your decision" in refused
    done = client.post(f"/hr/applications/{app_id}/override", data={
        "outcome": "SHORTLISTED", "designation": "ASSISTANT_PROFESSOR", "justification": "NET certificate seen in original."}).text
    assert "Overridden by HR" in done and "Justification: NET certificate seen in original." in done
    assert "Does not meet the minimum qualifications for any post." in done  # the finding is still shown beside it


def test_sending_back_from_the_page_reopens_the_fields_and_hides_the_old_assessment(client, engine, tmp_path):
    app_id, opening_id = _assessed_on_server(engine, tmp_path, _no_net())
    nothing = client.post(f"/hr/applications/{app_id}/return", data={"how": "correct"}).text
    assert "Tick at least one field to correct." in nothing
    back = client.post(f"/hr/applications/{app_id}/return", data={"how": "correct", "fields": ["net_set_status"]}).text
    assert "Fields to check (2)" in back and "An earlier assessment was sent back" in back and "Your decision" not in back
    client.post(f"/hr/applications/{app_id}/review", data={"f_net_set_status": "NET"})
    assert "Assess 1 read application(s)" in client.get(f"/hr/openings/{opening_id}").text
    again = client.post(f"/hr/openings/{opening_id}/assess").text
    assert "1 meet the post applied for" in again
    assert "Sent back" in client.get(f"/hr/applications/{app_id}").text  # the history of this gate


def test_the_excel_download_is_a_workbook_named_for_the_opening(client, engine, tmp_path):
    app_id, opening_id = _assessed_on_server(engine, tmp_path)
    r = client.get(f"/hr/openings/{opening_id}/export.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert f'OPN-{opening_id:05d}.xlsx' in r.headers["content-disposition"]
    assert load_workbook(io.BytesIO(r.content))["candidates"].max_row == 2
    assert "Download as Excel" in client.get(f"/hr/openings/{opening_id}").text
    assert client.get("/hr/openings/999999/export.xlsx").status_code == 404


# --- gaps found by walking through the application as HR would ----------------


def test_sending_back_offers_ready_ticked_the_fields_that_would_settle_what_is_open(session, tmp_path, client, engine):
    a = _assessed(session, tmp_path, _class_unstated(), school_id="SCH-008", discipline_group="ENGINEERING_TECHNOLOGY")
    assert a.status == states.MANUAL_REVIEW
    assert gate2.fields_that_would_settle(latest_evaluation(session, a.application_id)) == ["ug_marks_pct", "marks_pct"]
    settled = _assessed(session, tmp_path, _other_person())
    assert gate2.fields_that_would_settle(latest_evaluation(session, settled.application_id)) == []
    assert gate2.fields_that_would_settle(None) == []


@pytest.mark.parametrize("change, fields", [
    ({"net_set_status": ("SET", None)}, ["set_state"]),
    ({"marks_pct": (None, None), "cgpa": (8.4, "CGPA 8.4")}, ["marks_pct"]),
])
def test_each_kind_of_open_point_names_its_field(session, tmp_path, change, fields):
    r = _other_person()
    for name, (value, evidence) in change.items():
        f = getattr(r, name)
        f.value, f.evidence, f.confidence = value, evidence, 0.9 if value is not None else 0.0
    if "net_set_status" in change:
        r.net_set_status.evidence, r.set_state.value = "Cleared State Eligibility Test", None
    o = _opening(session, school_id="SCH-013", discipline_group="GENERAL")
    a = _application(session, tmp_path, category="General", differently_abled=False, study_leave_taken=False)
    a.opening_id, a.school_id = o.opening_id, o.school_id
    read_application(session, a, FakeProvider(script=[r]))
    if a.status == states.PENDING_REVIEW:  # the SET without a State stops at Gate 1 first; a person leaves it empty
        gate1.save_review(session, a, {}, {f.name for f in gate1.flagged_fields(session, a)})
    assess_application(session, a, TODAY)
    assert a.status == states.MANUAL_REVIEW
    assert gate2.fields_that_would_settle(latest_evaluation(session, a.application_id)) == fields


def test_withdrawing_after_a_decision_drops_the_letter_that_was_waiting(session, tmp_path):
    from backend import emails

    a = _assessed(session, tmp_path)
    gate2.approve(session, a)
    draft = emails.active_draft(session, a.application_id)
    emails.approve(session, draft)
    gate1.withdraw(session, a)
    assert a.status == states.WITHDRAWN and draft.status == "DISCARDED" and emails.active_draft(session, a.application_id) is None
    assert gate2.final_decision(session, a) is None


def test_a_field_can_be_reopened_before_assessment_from_the_page(client, engine, tmp_path):
    from sqlalchemy.orm import Session as _Session

    with _Session(engine) as s:
        seed_rules(s)
        o = _opening(s, school_id="SCH-013", discipline_group="GENERAL")
        a = _application(s, tmp_path, category="General", differently_abled=False, study_leave_taken=False)
        a.opening_id = o.opening_id
        read_application(s, a, FakeProvider(script=[_clean_result()]))
        s.commit()
        app_id, opening_id = a.application_id, o.opening_id
    page = client.get(f"/hr/applications/{app_id}").text
    assert "Something read wrongly? Correct a field before assessment" in page
    opened = client.post(f"/hr/applications/{app_id}/reopen", data={"fields": ["marks_pct"]}).text
    assert "Fields to check (1)" in opened and 'name="f_marks_pct"' in opened and 'value="68.4"' in opened
    client.post(f"/hr/applications/{app_id}/review", data={"f_marks_pct": "58.4"})
    assert client.get(f"/applications/{app_id}").json()["extracted"]["marks_pct"] == 58.4
    listing = client.get(f"/hr/openings/{opening_id}").text
    assert f"http://testserver/apply/{opening_id}" in listing  # the whole link, ready to copy
