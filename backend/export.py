"""The Excel download for one opening.

Laid out like the workbook the Reader stage produced, which HR already knows:
a `candidates` sheet with one row each, then one sheet per part of the record
(education, publications, seminars and workshops, teaching and skills,
experience, patents/awards/projects, guidance and memberships, contact
details). Three things are different:

* Rows are keyed by the application reference and listed in order of receipt.
  There is no rank and no home-made score: the system short-lists, it does
  not rank (cl. 4.1 Note).
* `candidates` leads with the engine's finding, its clause and page, and HR's
  decision; `assessment_checks` lists every requirement checked.
* The fields added in Phase 4 are there: authors and impact factor, the level
  of events and awards, funding amounts, subject levels, the listed discipline
  and the institution's tier.

Built from the database, in memory, and sent to the person who asked for it.
It holds personal data; nothing is written to disk here.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.detail_sheets import (
    NOT_FOUND,
    approx_years,
    degree_and_course,
    designation_text,
    division_text,
    duration_text,
    label,
    partial_date,
    tidy,
    title_text,
)
from app.excel_writer import EXCEL_DATE_FORMAT, SOURCE_CONVERTED, SOURCE_STATED, cgpa_to_percentage, date_as_stated
from backend import gate2, states
from backend.assessor_service import latest_evaluation
from backend.models import (
    Application,
    CandidateAchievement,
    CandidateEvent,
    CandidateExperience,
    CandidateGuidance,
    CandidateMembership,
    CandidatePersonalDetails,
    CandidatePublication,
    CandidateQualification,
    CandidateSkill,
    CandidateSubjectTaught,
    InstitutionMaster,
    JobOpening,
)
from llm.postprocess import format_degree, format_person_name, format_phd_evidence, format_phd_status

_RANKS = {"ASSISTANT_PROFESSOR": "Assistant Professor", "ASSOCIATE_PROFESSOR": "Associate Professor",
          "PROFESSOR": "Professor", "SENIOR_PROFESSOR": "Senior Professor"}
_FINDINGS = {"SHORTLISTED": "Meets the post applied for", "RE_CATEGORISED": "Meets a lower post",
             "NOT_ELIGIBLE": "Does not meet the minimum qualifications", "MANUAL_REVIEW": "For a person to decide"}
_ACTIONS = {"APPROVED": "Approved as assessed", "OVERRIDDEN": "Overridden", "DECIDED": "Decided by HR"}
_RESULTS = {"PASS": "Met", "FAIL": "Not met", "UNKNOWN": "Open"}
_SOURCES = {"WEB_FORM": "Application form", "MANUAL_UPLOAD": "HR upload", "EMAIL": "Email", "GOOGLE_FORM": "Google Form"}
_LEVELS = {"UG": "Bachelor's", "PG": "Master's", "PhD": "Ph.D.", "Post-Doc": "Post-Doc"}
_LEVEL_ORDER = {"UG": 0, "PG": 1, "PhD": 2, "Post-Doc": 3}
_GUIDANCE_LEVELS = {"PHD": "Ph.D.", "PG": "Master's", "UG": "Bachelor's", "OTHER": "Other"}

CANDIDATE_COLUMNS = (
    "reference", "candidate_name", "status", "applied_for",
    "finding", "meets_post", "reason", "clause", "page", "open_points",
    "hr_decision", "decided_post", "hr_justification", "decided_on",
    "highest_degree", "highest_degree_course_name", "masters_percentage", "masters_percentage_source", "masters_cgpa",
    "phd_status", "phd_course_name", "phd_award_date", "phd_regulation", "masters_award_date",
    "net_set_status", "set_state", "teaching_years_stated", "experience_counted_years",
    "publications_count", "publications_in_progress_count", "research_score_table_2", "shortlisting_score_table_3a",
    "category", "state", "differently_abled", "study_leave_taken", "fields_to_check",
    "rules_applied", "received", "source", "source_filename", "extraction_model",
)
KEY = ("reference", "candidate_name")
EDUCATION_COLUMNS = (*KEY, "level", "degree_and_course", "discipline_listed", "college", "university", "institution_tier",
                     "percentage", "cgpa", "division", "completed", "phd_status", "phd_thesis_title", "phd_guide",
                     "phd_registered", "check")
PUBLICATION_COLUMNS = (*KEY, "published_count", "in_progress_count", "no", "title", "type", "journal_or_conference", "year",
                       "status", "indexing", "authors", "first_author", "impact_factor", "check")
EVENT_COLUMNS = (*KEY, "no", "type", "title", "role", "level", "organiser", "duration", "year", "check")
TEACHING_COLUMNS = (*KEY, "teaching_years_stated", "subjects_count", "subjects_taught", "skills_count", "skills")
EXPERIENCE_COLUMNS = (*KEY, "teaching_years_stated", "no", "designation", "institution", "type", "from", "to", "approx_years",
                      "duration_as_stated", "held_while_studying", "check")
ACHIEVEMENT_COLUMNS = (*KEY, "no", "type", "title", "details", "level", "amount_as_stated", "amount_rupees", "year", "status", "check")
GUIDANCE_COLUMNS = (*KEY, "category", "level", "students", "detail", "check")
CONTACT_COLUMNS = (*KEY, "email", "phone", "state", "source_filename")
CHECK_COLUMNS = ("reference", "candidate_name", "post_checked", "requirement", "result", "what_was_compared", "clause", "page")

HEADER_FILL = PatternFill("solid", start_color="E8E8E8", end_color="E8E8E8")
NEEDS_PERSON_FILL = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
_WIDE = {
    "candidate_name": 28, "status": 44, "finding": 34, "reason": 60, "clause": 34, "open_points": 70, "hr_justification": 60,
    "highest_degree_course_name": 44, "phd_course_name": 44, "masters_percentage_source": 24, "fields_to_check": 46,
    "source_filename": 34, "degree_and_course": 40, "discipline_listed": 30, "college": 38, "university": 34, "division": 24,
    "phd_thesis_title": 50, "phd_guide": 26, "check": 30, "title": 66, "journal_or_conference": 48, "organiser": 44,
    "subjects_taught": 80, "skills": 70, "designation": 30, "institution": 48, "details": 50, "detail": 80, "email": 36,
    "requirement": 40, "what_was_compared": 90, "category": 24, "research_score_table_2": 24, "shortlisting_score_table_3a": 26,
    "experience_counted_years": 24, "rules_applied": 24, "received": 18, "decided_on": 18,
}
_WRAP = frozenset({"reason", "open_points", "hr_justification", "fields_to_check", "title", "subjects_taught", "skills",
                   "details", "detail", "what_was_compared", "phd_thesis_title"})


def _span(b: dict | None) -> Any:
    """A quantity that may be a range: a number when exact, text when not."""
    if not b:
        return None
    if b.get("high") is None:
        return f"at least {b['low']:g}"
    return b["low"] if b["low"] == b["high"] else f"{b['low']:g} to {b['high']:g}"


def _local(value: datetime | None) -> datetime | None:
    """Stored in UTC; shown in the server's local time, as a real date-time Excel can sort."""
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone().replace(tzinfo=None)


def _safe(value: Any) -> Any:
    """Text a spreadsheet will show as typed, never run as a formula."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _yes_no(value: bool | None) -> str | None:
    return None if value is None else ("Yes" if value else "No")


def _check(found: bool) -> str | None:
    return None if found else NOT_FOUND


def _reasons(reasons: list[str]) -> str | None:
    parts = []
    for reason in reasons or []:
        name, _, code = reason.partition(":")
        part = f"{name} ({code.replace('_', ' ')})" if code else name
        if part not in parts:
            parts.append(part)
    return "; ".join(parts) or None


def _course_name(e, qualifications: list[CandidateQualification]) -> str | None:
    """The highest degree with its course, in the house style; the fuller of the quote and the education row."""
    best = format_degree((e.evidence or {}).get("highest_degree"))
    degree = (best or "").split(" ")[0]
    for q in qualifications:
        if q.degree_level != e.highest_degree or not q.found_in_resume:
            continue
        candidate = format_degree(" ".join(p for p in (q.degree, q.discipline) if p))
        if candidate and (not degree or candidate.split(" ")[0] == degree) and len(candidate) > len(best or ""):
            best = candidate
    return best


def _phd_course_name(e, qualifications: list[CandidateQualification]) -> str | None:
    """ "Ph.D. <course>". The education row is used when it names the course: the quote behind the
    status often carries the university and dates as well, which do not belong in this column."""
    if e.phd_status in (None, "NOT_APPLICABLE"):
        return format_phd_evidence(e.phd_status, None)
    course = next((q.discipline for q in qualifications if q.degree_level == "PhD" and q.found_in_resume and q.discipline), None)
    if course:
        return format_degree(f"Ph.D. {course}")
    return format_phd_evidence(e.phd_status, (e.evidence or {}).get("phd_status"))


def opening_workbook(session: Session, opening: JobOpening, state_labels: dict[str, str]) -> bytes:
    now = datetime.now()
    tiers = {i.institution_id: i.tier for i in session.scalars(select(InstitutionMaster))}
    sheets: dict[str, tuple[tuple[str, ...], list[list[Any]]]] = {
        "candidates": (CANDIDATE_COLUMNS, []), "education": (EDUCATION_COLUMNS, []), "publications": (PUBLICATION_COLUMNS, []),
        "seminars_workshops": (EVENT_COLUMNS, []), "teaching_skills": (TEACHING_COLUMNS, []),
        "experience": (EXPERIENCE_COLUMNS, []), "patents_awards_projects": (ACHIEVEMENT_COLUMNS, []),
        "guidance_memberships": (GUIDANCE_COLUMNS, []), "contact_details": (CONTACT_COLUMNS, []),
        "assessment_checks": (CHECK_COLUMNS, []),
    }
    needs_person: list[int] = []  # rows of `candidates` to shade

    applications = session.scalars(
        select(Application).where(Application.opening_id == opening.opening_id).order_by(Application.application_id)
    ).all()
    for a in applications:
        def rows(table, key):
            return session.scalars(select(table).where(table.application_id == a.application_id).order_by(key)).all()

        personal = session.get(CandidatePersonalDetails, a.candidate_id)
        name = a.applicant_name or (personal.full_name if personal else None)
        key = [a.reference, name]
        e, v = a.extracted, latest_evaluation(session, a.application_id)
        details = (v.details or {}) if v and (a.status in gate2.AWAITING_HR or a.status == states.HR_APPROVED) else {}
        shown = v if details else None  # an assessment that was sent back is not reported as the finding
        decided = gate2.final_decision(session, a)
        clause, _, reason = (shown.failing_clause or "").partition(": ") if shown else ("", "", "")
        qualifications = rows(CandidateQualification, CandidateQualification.qualification_id)

        percentage = source = None
        if e is not None and e.marks_pct is not None:
            percentage, source = e.marks_pct, SOURCE_STATED
        elif e is not None and cgpa_to_percentage(e.cgpa) is not None:
            percentage, source = cgpa_to_percentage(e.cgpa), SOURCE_CONVERTED

        sheets["candidates"][1].append([
            a.reference, name, state_labels.get(a.status, a.status), _RANKS.get(a.applied_designation, a.applied_designation),
            _FINDINGS.get(shown.outcome, shown.outcome) if shown else None,
            _RANKS.get(shown.eligible_designation) if shown and shown.eligible_designation else None,
            reason or None, clause or None, shown.failing_clause_page if shown else None,
            "; ".join(details.get("open_points") or []) or None,
            _ACTIONS.get(decided.action) if decided else None,
            (_RANKS.get(decided.final_designation) or ("None" if decided.final_outcome == "NOT_ELIGIBLE" else None)) if decided else None,
            decided.justification if decided else None, _local(decided.decided_at) if decided else None,
            e.highest_degree if e else None, _course_name(e, qualifications) if e else None,
            percentage, source, e.cgpa if e else None,
            format_phd_status(e.phd_status) if e else None,
            _phd_course_name(e, qualifications) if e else None,
            date_as_stated(e.phd_award_date, e.phd_award_date_precision) if e else None, e.phd_regulation if e else None,
            date_as_stated(e.masters_award_date, e.masters_award_date_precision) if e else None,
            e.net_set_status if e else None, e.set_state if e else None, e.teaching_years if e else None,
            _span(details.get("experience_years")),
            e.publications_count if e else None, e.publications_in_progress_count if e else None,
            _span(details.get("research_score")), _span(details.get("shortlist_score")),
            a.category, personal.state if personal else None, _yes_no(a.differently_abled), _yes_no(a.study_leave_taken),
            _reasons(e.review_reasons) if e and a.status == states.PENDING_REVIEW else None,
            details.get("rule_version"), _local(a.created_at), _SOURCES.get(a.resume_source, a.resume_source),
            a.resume_filename, e.model_used if e else None,
        ])
        if a.status in (states.PENDING_REVIEW, states.MANUAL_REVIEW, states.FAILED):
            needs_person.append(len(sheets["candidates"][1]) + 1)

        for q in sorted(qualifications, key=lambda q: _LEVEL_ORDER.get(q.degree_level, 9)):
            is_phd = q.degree_level == "PhD"
            sheets["education"][1].append([
                *key, _LEVELS.get(q.degree_level, q.degree_level), degree_and_course(q.degree, q.discipline), q.discipline_listed,
                title_text(q.college_name), title_text(q.university_name), tiers.get(q.institution_id),
                q.marks_pct, q.cgpa, division_text(q.division), partial_date(q.completion_stated),
                format_phd_status(q.phd_status) if is_phd else None, title_text(q.thesis_title) if is_phd else None,
                format_person_name(tidy(q.guide)) if is_phd else None, partial_date(q.registration_stated) if is_phd else None,
                _check(q.found_in_resume),
            ])

        published = e.publications_count if e else None
        in_progress = e.publications_in_progress_count if e else None
        for i, p in enumerate(rows(CandidatePublication, CandidatePublication.publication_id), start=1):
            sheets["publications"][1].append([
                *key, published, in_progress, i, title_text(p.title), label(p.publication_type), title_text(p.venue_name),
                p.year, label(p.status), tidy(p.indexing), p.author_count, _yes_no(p.is_first_author), p.impact_factor,
                _check(p.found_in_resume),
            ])
        for i, x in enumerate(rows(CandidateEvent, CandidateEvent.id), start=1):
            sheets["seminars_workshops"][1].append([
                *key, i, label(x.kind), title_text(x.title), label(x.role), label(x.level), title_text(x.organiser),
                duration_text(x.duration_stated), x.year, _check(x.found_in_resume),
            ])

        subjects = rows(CandidateSubjectTaught, CandidateSubjectTaught.id)
        skills = rows(CandidateSkill, CandidateSkill.id)
        if e is not None:
            sheets["teaching_skills"][1].append([
                *key, e.teaching_years, len(subjects),
                "; ".join(tidy(s.subject_name) + (f" ({s.course_level})" if s.course_level else "") for s in subjects) or None,
                len(skills), "; ".join(tidy(s.skill_name) for s in skills) or None,
            ])
        for i, x in enumerate(rows(CandidateExperience, CandidateExperience.experience_id), start=1):
            sheets["experience"][1].append([
                *key, e.teaching_years if e else None, i, designation_text(x.designation_held), title_text(x.employer_name),
                label(x.experience_type), partial_date(x.start_stated), "Present" if x.is_current else partial_date(x.end_stated),
                approx_years(x.start_stated, "PRESENT" if x.is_current else x.end_stated, now), duration_text(x.duration_stated),
                _yes_no(x.concurrent_with_study), _check(x.found_in_resume),
            ])
        for i, x in enumerate(rows(CandidateAchievement, CandidateAchievement.id), start=1):
            sheets["patents_awards_projects"][1].append([
                *key, i, label(x.kind), title_text(x.title), tidy(x.details), label(x.level), x.amount_stated, x.amount_inr,
                x.year, tidy(x.status), _check(x.found_in_resume),
            ])
        for g in rows(CandidateGuidance, CandidateGuidance.id):
            sheets["guidance_memberships"][1].append([
                *key, "Research guidance", _GUIDANCE_LEVELS.get(g.level, g.level), g.student_count, tidy(g.description),
                _check(g.found_in_resume),
            ])
        for m in rows(CandidateMembership, CandidateMembership.id):
            sheets["guidance_memberships"][1].append([*key, "Professional membership", None, None, tidy(m.membership), None])
        sheets["contact_details"][1].append([
            *key, a.applicant_email or (personal.contact_email if personal else None),
            a.applicant_phone or (personal.contact_phone if personal else None), personal.state if personal else None,
            a.resume_filename,
        ])
        for rank in details.get("ranks") or []:
            for c in rank["checks"]:
                sheets["assessment_checks"][1].append([
                    *key, _RANKS.get(rank["designation"], rank["designation"]), c["label"], _RESULTS.get(c["result"], c["result"]),
                    c["detail"], c["clause"], c["page"],
                ])

    wb = Workbook()
    wb.remove(wb.active)
    for title, (columns, data) in sheets.items():
        ws = wb.create_sheet(title)
        ws.append(columns)
        for row in data:
            ws.append([_safe(v) for v in row])
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True), HEADER_FILL
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        for i, column in enumerate(columns, start=1):
            letter = get_column_letter(i)
            ws.column_dimensions[letter].width = _WIDE.get(column, 16)
            for cell in ws[letter][1:]:
                if isinstance(cell.value, datetime):
                    cell.number_format = "dd-mm-yyyy hh:mm"
                elif hasattr(cell.value, "year"):
                    cell.number_format = EXCEL_DATE_FORMAT
                if column in _WRAP:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row_number in needs_person:
        for cell in wb["candidates"][row_number]:
            cell.fill = NEEDS_PERSON_FILL

    note = wb.create_sheet("how_to_read")
    for line in (
        f"{opening.reference}: {_RANKS.get(opening.designation, opening.designation)}, {opening.school.name}",
        f"Downloaded {now.strftime('%d-%m-%Y %H:%M')}.",
        "",
        "A short-listing aid. Selection is made by the Selection Committee at interview (UGC Regulations, 2018, clause 4.1 Note).",
        "Rows are in order of receipt. Nothing here is a ranking, and the short-listing score must not be used as one.",
        "",
        "candidates: one row per application. 'finding' is what the rules give on the record as read, with its clause and page;",
        "'hr_decision' is what a person decided. A shaded row is waiting for a person: fields to check, a point the rules could not settle, or a file that could not be read.",
        "The other sheets hold each part of the record, one row per item, keyed by the same reference.",
        "assessment_checks: every requirement checked for every post looked at, with the figures compared.",
        "",
        "'check' says 'Not found in resume text - verify' where an item the reader returned could not be located in the resume.",
        "masters_percentage_source says whether the percentage was stated on the resume or converted from a CGPA for display. A converted figure is never used to decide anything.",
        "Scores and experience are ranges where the resume or the gazette leaves something open.",
        "experience_counted_years counts teaching, research and industry posts under an AICTE rule, and teaching and research posts under a UGC rule.",
        "An empty cell means the resume did not state it. Nothing is filled in by guesswork.",
        "",
        "This file holds personal data. Keep it as you would the applications themselves.",
    ):
        note.append([line])
    note.column_dimensions["A"].width = 150
    note["A1"].font = Font(bold=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
